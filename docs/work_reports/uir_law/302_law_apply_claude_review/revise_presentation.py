"""Revise a copy of the KAIC oral deck without rerunning or altering LAW-KR-200."""
from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
PAPERS = HERE.parents[2] / "papers/_47_2_UIR_Law"
SOURCE = PAPERS / "2026-09-14_#0021_UIR_KAIC2026_Oral_Presentation.pptx"
OUTPUT = PAPERS / "2026-09-19_#0024_UIR_KAIC2026_Oral_Presentation_REVIEWED.pptx"


def rewrite(prs: Presentation, slide: int, shape: int, paragraph: int, value: str, size: float | None = None) -> None:
    target = prs.slides[slide - 1].shapes[shape].text_frame.paragraphs[paragraph]
    if not target.runs:
        raise ValueError((slide, shape, paragraph))
    run_props = copy.deepcopy(target.runs[0]._r.rPr)
    # Paragraph.text converts vertical tabs into native DrawingML <a:br/> runs.
    # Run.text would escape them as literal _x000B_, corrupting slide layout.
    target.text = value
    for run in target.runs:
        old_props = run._r.rPr
        if old_props is not None:
            run._r.remove(old_props)
        if run_props is not None:
            run._r.insert(0, copy.deepcopy(run_props))
        if size is not None:
            run.font.size = Pt(size)


def note(prs: Presentation, slide: int, value: str) -> None:
    prs.slides[slide - 1].notes_slide.notes_text_frame.text = value


def build() -> Path:
    prs = Presentation(SOURCE)
    assert len(prs.slides) == 16
    rewrite(prs, 1, 1, 1, "UIR: 환각 저항형 법률 검색을 위한\v범용 중간 표현")

    rewrite(prs, 3, 4, 1,
        "• 텍스트 프롬프트에 의도와 미검증 개체가 함께 들어옴\v"
        "  - 사용자의 사건번호 주장을 사실 권위로 승격할 위험\v\v"
        "• 비정확한 검색의 전파\v"
        "  - 본 실험의 lexical top-3 검색은 다른 실존 사건을 반환할 수 있음\v\v"
        "• 생성 출력의 identity 오류\v"
        "  - 다른 사건의 metadata를 원래 질의에 대한 답으로 제시 가능")
    rewrite(prs, 3, 5, 0,
        "평가된 실패 경로: 질의 속 사건번호 주장을 검증 없이 사실 권위로 취급하면 다른 사건의 identity가 전파될 수 있다.")
    rewrite(prs, 4, 2, 1,
        "• 구현: 결정적 lexical top-3 검색 후 문맥을 모델에 전달\v\v"
        "• 검증 없이 다른 실존 사건으로 답할 수 있음\v\v"
        "• LAW-KR-200 관찰:\v  AULCR 96.39% (80/83), verified coverage 3%")
    rewrite(prs, 4, 3, 1,
        "• 구현: 생성 전 사건번호 존재 여부만 확인\v\v"
        "• 실존 번호 + 다른 사건명인 D 질의는 결합 검증 없음\v\v"
        "• LAW-KR-200 관찰:\v"
        "  AULCR 69.01% (98/142)\v"
        "  D safe outcome 19/50 (38%)")
    rewrite(prs, 4, 4, 1,
        "• C4: 모델 제어 lookup tool calling; valid FRR 100%\v\v"
        "• C5: naive retrieval + 사후 citation/name/evidence allow-list\v"
        "  AULCR 75.61%, valid FRR 84%\v\v"
        "• 두 방법 모두 평가 범위에서 safety–utility 균형에 한계")
    rewrite(prs, 4, 5, 0,
        "LAW-KR-200 관찰: 단순 검색은 비근거 수용이 높았고, 도구·사후 검증 방식은 유효 질의 거절이 높았다.")
    rewrite(prs, 5, 3, 1,
        "• 자연어 프롬프트: 사용자 의도와 맥락 전달\v\v"
        "• 본 평가의 권위 경계: frozen 국가법령정보 공동활용 snapshot\v"
        "  - 사건번호 존재와 사건명 결합 확인\v"
        "  - 미확인 개체를 사실로 승격하지 않음")

    rewrite(prs, 6, 4, 1,
        "• 생성 전 frozen 국가법령정보 공동활용 snapshot과 대조\v\v"
        "• 두 가지 확인:\v"
        "  1) exact 사건번호 match 존재 여부\v"
        "  2) 사건번호 ↔ 사건명 일치 여부\v\v"
        "• VERIFIED → 검색·생성\v"
        "• NOT_FOUND / MISMATCH → 기권")
    rewrite(prs, 6, 7, 0,
        "핵심: 평가된 NOT_FOUND/MISMATCH 요청은 모델 호출 전에 기권하여 미검증 개체의 자유 생성을 막았다.")
    rewrite(prs, 6, 6, 1,
        "• 명시적 기권 (Safe Abstention)\v"
        "• 평가된 경로에서 model_invoked=false\v"
        "• 비존재/불일치 유형을 응답에 표시\v"
        "→ 미검증 개체에 대한 생성 단계 없음")
    rewrite(prs, 7, 3, 1,
        "• 모델 호출 전 frozen authoritative snapshot과 사건번호 binding\v"
        "• 사건번호 존재와 사건명 결합 일치성 검증\v"
        "• 미확인 번호·mismatch는 평가된 경로에서 생성 전 기권")
    rewrite(prs, 7, 4, 1,
        "• NOT_FOUND/MISMATCH 시 명시적 기권\v"
        "• 이 경로에서는 LLM 호출 없음 (model_invoked=false)\v"
        "• 실패 유형을 숨기지 않는 안전한 응답")
    rewrite(prs, 7, 5, 1,
        "• 생성 출력의 citation, case name, source ID를 재확인\v"
        "• 다른 사건의 identity 삽입을 출력 단계에서 검출\v"
        "• 결합된 evidence ID가 있는 출력만 수용")
    rewrite(prs, 7, 6, 0,
        "인과 사슬: 의미 정규화 → 개체 검증 → 근거 결합 법률 검색 (C-label은 내부 baseline 식별자)")
    rewrite(prs, 8, 3, 1,
        "• 한국어–영어 질의 100쌍 (동일 underlying entity)\v\v"
        "• Dev 40 / Test 200 분리·동결\v"
        "  - parser 수정은 dev에서만\v\v"
        "• 런타임은 opaque ID와 query만 접근\v"
        "  - gold pre-generation access: 0건\v\v"
        "• Phi-3.5-mini-instruct (3.8B, Q4_0)\v"
        "  - seed=20260911, temperature=0, SHA-256 기록")
    rewrite(prs, 8, 4, 0,
        "평가 범위: 동결된 한국 판례 snapshot의 사건번호·사건명 identity 검증이며 실체적 법리 정확도는 평가하지 않았다.")

    rewrite(prs, 9, 2, 1,
        "1. AULCR = Accepted Unsupported Legal Claim Rate ↓\v"
        "   - ANSWER/CORRECT 중 근거 없는 법률 주장을 포함한 출력 / accepted outputs\v\v"
        "2. Invalid Citation FAR ↓\v"
        "   - Category C hard negative의 accepted / 50; D는 별도 binding 평가\v\v"
        "3. Verified Coverage ↑\v"
        "   - 유효 A+B에서 source-bound identity answer / 100\v\v"
        "4. FRR ↓\v"
        "   - 유효 A+B의 REJECT/CLARIFY / 100")
    rewrite(prs, 9, 3, 1,
        "• Direct  [C0]  — 검색·검증 없음\v\v"
        "• Naive RAG  [C1]  — 결정적 lexical top-3; entity 검증 없음\v\v"
        "• Existence Check  [C2]  — 번호 존재만 확인\v\v"
        "• Tool Agent  [C4]  — model-controlled lookup tool calling\v\v"
        "• Guardrail  [C5]  — naive retrieval + 사후 citation/name/evidence allow-list\v\v"
        "• UIR  [C8]  — typed binding + 출력 계약")

    for shape_index in (4, 5, 6):
        card = prs.slides[9].shapes[shape_index]
        card.top = Inches(4.85)
        card.height = Inches(1.30)
        card.text_frame.margin_top = Inches(0.05)
        card.text_frame.margin_bottom = Inches(0.05)
    rewrite(prs, 10, 4, 0, "UIR: 0건 관찰", size=11)
    rewrite(prs, 10, 4, 1,
        "AULCR = Accepted Unsupported Legal Claim Rate\v"
        "0/100 (Wilson 95% CI [0.00%, 3.70%])\v"
        "Valid identity coverage 100/100; FRR 0/100", size=9.0)
    rewrite(prs, 10, 5, 0, "Tool Agent: 0%의 해석", size=11)
    rewrite(prs, 10, 5, 1,
        "Accepted 1건; AULCR 0/1\v"
        "Valid FRR 100/100 (100%)\v"
        "안전성만으로 유용성 주장 불가", size=9.5)
    rewrite(prs, 10, 6, 0, "Existence Check: 결합 한계", size=11)
    rewrite(prs, 10, 6, 1,
        "Invalid FAR (C) 0/50\v"
        "D safe outcome 19/50 (38%)\v"
        "AULCR 98/142 (69.01%)", size=9.5)
    rewrite(prs, 10, 7, 0,
        "LAW-KR-200: UIR의 accepted unsupported claim 0/100 관찰; valid identity coverage 100/100, FRR 0/100.")

    rewrite(prs, 11, 4, 1,
        "• 평가된 pipeline에서 safety와 valid coverage를 함께 비교\v"
        "• UIR 관찰값: AULCR 0/100, valid identity coverage 100/100\v"
        "• Tool Agent는 AULCR 0/1이나 valid FRR 100/100")
    rewrite(prs, 11, 4, 0, "시각화 시사점: 평가 표본에서의 안전성–유효 커버리지 비교")
    rewrite(prs, 12, 3, 1,
        "• SELECT EXISTS는 사건번호만 확인\v\v"
        "• Category D: 실존 번호 + 다른 사건명 50건\v"
        "  - Existence Check safe outcome 19/50 (38%)\v"
        "  - UIR safe outcome 50/50 (100%)\v\v"
        "• D safe outcome = 거절·명확화 또는 검증된 개체로 정정\v"
        "  (FRR은 유효 A/B에만 적용)")
    rewrite(prs, 12, 2, 1,
        "• 동일 200 질의의 paired exact McNemar 검증\v"
        "  (Holm-Bonferroni 다중비교 보정)\v\v"
        "• vs Naive RAG: p = 1.65 × 10⁻²⁴\v"
        "  (Holm p = 3.31 × 10⁻²⁴; 실패 80건 vs 0건)\v\v"
        "• vs Existence Check: p = 6.31 × 10⁻³⁰\v"
        "  (Holm p = 1.89 × 10⁻²⁹; 실패 98건 vs 0건)\v\v"
        "• vs Guardrail: p = 9.31 × 10⁻¹⁰\v"
        "  (identity 실패 31건 vs 0건)\v\v"
        "→ 동결 표본의 identity 오류에서 통계적 차이 관찰")
    rewrite(prs, 13, 2, 1,
        "• Naive RAG: 요청 2026두30951 대신 다른 사건 2026두31165를 ANSWER로 출력\v"
        "  - frozen raw 출력은 max-token 종료; 완결된 법리 요약으로 해석하지 않음\v\v"
        "• UIR: NOT_FOUND로 모델 호출 전 REJECT (model_invoked=false)\v"
        "  - 공식 HTTP 200 검색 응답 내 exact 사건번호 부재를 동결·검증")
    rewrite(prs, 13, 3, 1,
        "• Existence Check: 번호가 존재해도 사건명 불일치에 취약\v"
        "  - D safe outcome 19/50; 그 외 31건은 unsupported ANSWER\v\v"
        "• UIR: 사건번호–사건명 binding을 생성 전에 확인\v"
        "  - D safe outcome 50/50, 전건 REJECT 관찰")
    rewrite(prs, 13, 4, 0,
        "정성 사례: frozen snapshot의 exact-match와 identity binding이 미검증 개체의 모델 유입을 막는 과정을 보여준다.")
    rewrite(prs, 14, 3, 1,
        "• Canonical UIR equivalence: 100/100 한·영 쌍 관찰\v\v"
        "• KO/EN valid identity coverage: 각각 100%\v"
        "• KO/EN accepted unsupported incidence: 각각 0%\v\v"
        "• 다른 언어나 법역으로의 일반화는 평가하지 않음")
    rewrite(prs, 14, 4, 1,
        "• 언어별 frontend와 공통 typed UIR/backend 분리\v"
        "• 본 실험은 한국어·영어만 평가\v"
        "• 추가 언어의 parser 정확도와 backend 재사용성은 향후 검증")
    rewrite(prs, 14, 5, 0,
        "관찰 범위: 동결된 한국 대법원 사건번호의 한·영 질의 100쌍에서 동일 canonical UIR을 얻었다.")
    rewrite(prs, 15, 2, 1,
        "1. 사건번호–사건명 identity metadata 평가\v"
        "   - 실체적 법리 추론·판결 요지 정확도는 미평가\v\v"
        "2. 0건 관찰 ≠ 0 risk\v"
        "   - UIR AULCR 0/100; Wilson 95% CI [0.00%, 3.70%]\v\v"
        "3. 단일 frozen 국가법령정보 공동활용 snapshot\v"
        "   - 정보원 불일치·source outage·schema 변경 미평가\v\v"
        "4. 로컬 Phi-3.5-mini-instruct 3.8B의 통제 결과")
    rewrite(prs, 15, 3, 1,
        "1. CourtListener 기반 LAW-US-200 독립 replication\v\v"
        "2. Source unavailability와 다중 권위 충돌의 fail-closed 동작 평가\v\v"
        "3. Typed UIR과 structured legal database / legacy RDB mapping\v\v"
        "4. Agent–RDB orchestration의 범위 확장")
    rewrite(prs, 16, 2, 1,
        "• 한·영 질의를 공통 typed UIR로 정규화\v"
        "• 생성 전 frozen 국가법령정보 공동활용 snapshot과 identity binding\v"
        "• 검증 실패 시 평가된 경로에서 모델 호출 없이 기권")
    rewrite(prs, 16, 3, 1,
        "• AULCR 0/100 관찰 (Wilson 95% CI [0.00%, 3.70%])\v"
        "• Valid source-bound identity coverage 100/100, FRR 0/100\v"
        "• 실체적 법률 추론의 일반적 안전성을 뜻하지 않음")
    rewrite(prs, 16, 3, 0, "3. 실증적 관찰\vLAW-KR-200")

    note(prs, 3,
        "사건번호가 포함된 질의는 사용자의 의도와 미검증 개체 주장을 함께 전달합니다. "
        "LAW-KR-200의 lexical top-3 baseline에서는 다른 실존 사건이 검색되어 원래 질의의 identity로 잘못 제시될 수 있었습니다. "
        "이는 평가된 시스템 경로의 관찰이며 모든 RAG 구현에 대한 일반적 주장이나 법리 내용 평가가 아닙니다.")
    note(prs, 1,
        "안녕하십니까. 'UIR: 환각 저항형 법률 검색을 위한 범용 중간 표현'을 발표할 아주대학교 컴퓨터공학과 박성수입니다. "
        "본 연구는 법률 정보 검색에서 미검증 판례 identity가 답변으로 전파되는 문제를 다룹니다. "
        "컴파일러 중간 표현에서 착안한 생성 전 검증 구조와 LAW-KR-200의 통제된 평가 결과를 말씀드리겠습니다.")
    note(prs, 4,
        "Naive RAG는 결정적 lexical top-3 검색을 사용했고 AULCR 80/83, valid identity coverage 3/100을 기록했습니다. "
        "번호 존재만 확인한 C2의 Category D safe outcome은 19/50이며, 나머지 31건은 unsupported ANSWER였습니다. "
        "C4는 모델 제어 lookup tool calling으로 accepted output 1건, valid FRR 100/100이었습니다. "
        "C5는 naive retrieval 뒤 citation/name/evidence allow-list를 적용했으며 valid FRR 84/100이었습니다. "
        "이 수치는 각 구현과 동결 표본에 한정됩니다.")
    note(prs, 5,
        "사용자가 입력한 사건번호는 그 자체로 검증된 사실이 아닙니다. "
        "본 연구는 frozen 국가법령정보 공동활용 snapshot을 평가의 권위 경계로 삼아 사건번호와 사건명의 일치를 확인했습니다. "
        "단일 정보원의 오류, 부재, 다른 권위와의 충돌은 별도 평가가 필요합니다.")

    note(prs, 6,
        "이 실험의 공식 응답은 dataset construction 때 수집되어 frozen snapshot으로 실행됩니다. "
        "한국 API hard negative는 HTTP 200 목록 응답에서 exact 사건번호가 없을 때 normalized NOT_FOUND로 기록했습니다. "
        "UIR은 사건번호와 사건명 binding을 생성 전에 검증하며, 이 평가의 NOT_FOUND/MISMATCH 경로에서는 모델을 호출하지 않았습니다.")
    note(prs, 7,
        "실시간 API 조회가 아닙니다. 생성 전 frozen authoritative snapshot에서 typed entity를 결합합니다. "
        "출력 계약은 citation, case name, source ID의 일치를 확인합니다. 미평가된 장애·다중 정보원 상황은 안전성 보장 대상으로 주장하지 않습니다.")
    note(prs, 8,
        "LAW-KR-200은 한·영 100쌍의 질의로 이루어졌고 A/B/C/D가 각 50건입니다. "
        "Dev 40건에서 parser를 조정한 뒤 Test 200건의 질의와 source snapshot을 동결했습니다. "
        "평가 런타임에는 opaque ID와 query만 전달됐고 사전 gold 접근은 0건으로 기록됐습니다. "
        "따라서 여기서의 결론은 이 통제된 identity 벤치마크 범위에 한정합니다.")
    note(prs, 9,
        "6개 pipeline은 C0, C1, C2, C4, C5, C8이며 C-label은 내부 harness 식별자입니다. "
        "AULCR 분모는 ANSWER/CORRECT로 수용된 출력 수입니다. Invalid Citation FAR는 Category C의 accepted/50이고 D는 별도 binding safe outcome입니다. "
        "C1은 deterministic lexical top-3, C4는 model-controlled lookup, C5는 naive retrieval 뒤 사후 allow-list입니다.")
    note(prs, 10,
        "UIR은 accepted unsupported legal claim 0/100이 관찰되었습니다. Wilson 95% CI는 0.00%–3.70%로, 0 risk를 뜻하지 않습니다. "
        "Valid source-bound identity coverage는 100/100, false rejection은 0/100입니다. "
        "Tool Agent는 accepted output이 1건이며 AULCR 0/1이지만, valid 100건을 모두 거절했습니다.")
    note(prs, 11,
        "산점도는 frozen LAW-KR-200의 accepted unsupported claim rate와 유효 A/B의 source-bound identity coverage를 비교합니다. "
        "UIR은 이 표본에서 AULCR 0/100과 coverage 100/100이었지만 0 risk를 뜻하지 않습니다. "
        "C4는 accepted 1건에서 AULCR 0/1, valid FRR 100/100이었고 C5는 valid FRR 84/100으로 정상 질의 16건은 거절하지 않았습니다.")
    note(prs, 12,
        "Paired exact McNemar 검정과 Holm 보정은 이 동결 표본의 identity 실패 차이를 비교한 결과입니다. "
        "Category D는 FRR 대상이 아닙니다. Safe outcome은 명시적 거절·명확화 또는 source-bound correction이며, "
        "Existence Check 19/50, UIR 50/50이 관찰됐습니다. A/B에만 valid verified coverage와 false rejection을 적용합니다.")
    note(prs, 13,
        "실제 hard negative 2026두30951에 대해 Naive RAG raw 출력은 다른 사건 2026두31165를 ANSWER로 제시했으나 max-token에서 종료되었습니다. "
        "UIR은 frozen source의 exact-match 부재를 확인해 모델 호출 전에 REJECT했습니다. "
        "이 사례는 identity metadata 실패의 예시이며 판결 요지 정확도 비교가 아닙니다.")
    note(prs, 14,
        "동결된 한국 대법원 판례 질의 100개 한·영 쌍에서 canonical UIR 일치 100/100을 관찰했습니다. "
        "이 표본의 KO/EN valid identity coverage는 각각 100%이고 accepted unsupported incidence는 각각 0%였습니다. "
        "다른 언어와 법역의 parser 정확도 또는 backend 재사용성은 아직 평가하지 않았습니다.")
    note(prs, 15,
        "이 평가는 사건번호·사건명 identity metadata에 집중하며 판결의 실체적 법리 추론과 요지 정확도는 평가하지 않았습니다. "
        "AULCR 0/100의 Wilson 95% 상한은 3.70%입니다. 단일 frozen 정보원과 로컬 3.8B 모델을 사용했으므로 "
        "source outage, schema 변경, 다중 정보원 충돌 및 다른 법역의 일반화는 향후 과제입니다. "
        "CourtListener LAW-US-200은 별도의 독립 replication 대상으로 남아 있습니다.")
    note(prs, 16,
        "UIR은 이 동결 LAW-KR-200 평가에서 accepted unsupported legal claim 0/100, "
        "valid source-bound identity coverage 100/100, FRR 0/100을 관찰했습니다. "
        "이는 생성 전 identity binding과 기권 경로의 결과이며 실체적 법리 추론의 일반적 안전성을 입증하지 않습니다. "
        "경청해 주셔서 감사합니다.")

    # Re-read after serialization in QA; no source artifact is overwritten.
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
